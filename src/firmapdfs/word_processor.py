"""
Word Document Processor
=======================

Handles Word document processing including:
- Opening and saving documents
- Setting date fields
- Modifying checkboxes (content controls and legacy form fields)

Supports two engines:
- 'docx': python-docx library (cross-platform, limited control support)
- 'com': Windows COM automation via pywin32 (Windows only, full support)
"""

import re
import shutil
from abc import ABC, abstractmethod
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from dataclasses import dataclass, field

from .logger import get_logger
from .utils import format_date, ensure_unique_filename

logger = get_logger(__name__)


# =============================================================================
# Data Classes
# =============================================================================

@dataclass
class FieldInfo:
    """Information about a document field."""
    name: str
    field_type: str  # "content_control", "form_field", "bookmark", "text"
    value: Any = None
    tag: str = ""
    title: str = ""


@dataclass
class CheckboxInfo:
    """Information about a checkbox control."""
    name: str
    checked: bool
    control_type: str  # "content_control", "form_field", "legacy"
    tag: str = ""
    title: str = ""


@dataclass
class ProcessingResult:
    """Result of document processing."""
    success: bool
    input_path: Path
    output_path: Optional[Path] = None
    error_message: str = ""
    fields_modified: List[str] = field(default_factory=list)
    checkboxes_modified: List[str] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)


# =============================================================================
# Abstract Base Engine
# =============================================================================

class WordEngine(ABC):
    """
    Abstract base class for Word document processing engines.

    Provides a common interface for both python-docx and COM implementations.
    """

    def __init__(self, config: Any = None):
        """
        Initialize the engine.

        Args:
            config: WordConfig object with settings
        """
        self.config = config
        self._document = None
        self._document_path: Optional[Path] = None

    @abstractmethod
    def open(self, file_path: Path) -> bool:
        """
        Open a Word document.

        Args:
            file_path: Path to the .docx file

        Returns:
            True if successful, False otherwise
        """
        pass

    @abstractmethod
    def save(self, output_path: Optional[Path] = None) -> bool:
        """
        Save the document.

        Args:
            output_path: Path to save to (optional, saves in place if None)

        Returns:
            True if successful, False otherwise
        """
        pass

    @abstractmethod
    def close(self) -> None:
        """Close the document and release resources."""
        pass

    @abstractmethod
    def get_text_content(self) -> str:
        """
        Get the full text content of the document.

        Returns:
            Document text as a single string
        """
        pass

    @abstractmethod
    def find_checkboxes(self) -> List[CheckboxInfo]:
        """
        Find all checkbox controls in the document.

        Returns:
            List of CheckboxInfo objects
        """
        pass

    @abstractmethod
    def set_checkbox(self, name: str, checked: bool) -> bool:
        """
        Set a checkbox state by name or tag.

        Args:
            name: Checkbox name, tag, or partial text
            checked: True to check, False to uncheck

        Returns:
            True if found and modified, False otherwise
        """
        pass

    @abstractmethod
    def find_date_fields(self) -> List[FieldInfo]:
        """
        Find date fields in the document.

        Returns:
            List of FieldInfo objects for date fields
        """
        pass

    @abstractmethod
    def set_date_field(self, name: str, date: datetime) -> bool:
        """
        Set a date field value.

        Args:
            name: Field name or tag
            date: Date to set

        Returns:
            True if found and modified, False otherwise
        """
        pass

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        self.close()
        return False


# =============================================================================
# Python-docx Engine (Cross-platform)
# =============================================================================

class DocxEngine(WordEngine):
    """
    Word processing engine using python-docx library.

    Cross-platform but has limited support for content controls.
    Best for simple documents without complex form fields.
    """

    def __init__(self, config: Any = None):
        super().__init__(config)
        self._docx = None

    def open(self, file_path: Path) -> bool:
        """Open a Word document using python-docx."""
        try:
            from docx import Document

            self._document_path = Path(file_path)
            self._docx = Document(str(file_path))
            self._document = self._docx

            logger.debug(f"Opened document with python-docx: {file_path}")
            return True

        except Exception as e:
            logger.error(f"Failed to open document with python-docx: {e}")
            return False

    def save(self, output_path: Optional[Path] = None) -> bool:
        """Save the document."""
        if self._docx is None:
            logger.error("No document open to save")
            return False

        try:
            save_path = output_path or self._document_path
            self._docx.save(str(save_path))
            logger.debug(f"Saved document: {save_path}")
            return True

        except Exception as e:
            logger.error(f"Failed to save document: {e}")
            return False

    def close(self) -> None:
        """Close the document."""
        self._docx = None
        self._document = None
        self._document_path = None

    def get_text_content(self) -> str:
        """Get document text content."""
        if self._docx is None:
            return ""

        paragraphs = []
        for para in self._docx.paragraphs:
            paragraphs.append(para.text)

        # Also get text from tables
        for table in self._docx.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)

        return "\n".join(paragraphs)

    def find_checkboxes(self) -> List[CheckboxInfo]:
        """
        Find checkboxes in the document.

        Note: python-docx has limited native support for content controls.
        This implementation searches the underlying XML.
        """
        checkboxes = []

        if self._docx is None:
            return checkboxes

        try:
            # Access the document's XML to find content controls
            from docx.oxml.ns import qn

            # Search for sdt (Structured Document Tag) elements with checkbox
            for sdt in self._docx.element.iter(qn('w:sdt')):
                checkbox_info = self._parse_sdt_checkbox(sdt)
                if checkbox_info:
                    checkboxes.append(checkbox_info)

            # Also search for legacy form field checkboxes
            for ffdata in self._docx.element.iter(qn('w:ffData')):
                checkbox_info = self._parse_legacy_checkbox(ffdata)
                if checkbox_info:
                    checkboxes.append(checkbox_info)

        except Exception as e:
            logger.warning(f"Error finding checkboxes: {e}")

        return checkboxes

    def _parse_sdt_checkbox(self, sdt) -> Optional[CheckboxInfo]:
        """Parse a structured document tag for checkbox info."""
        try:
            from docx.oxml.ns import qn

            # Get SDT properties
            sdt_pr = sdt.find(qn('w:sdtPr'))
            if sdt_pr is None:
                return None

            # Check if it's a checkbox (w14:checkbox)
            checkbox = sdt_pr.find(qn('w14:checkbox'))
            if checkbox is None:
                # Also check for older format
                checkbox = sdt_pr.find('{http://schemas.microsoft.com/office/word/2010/wordml}checkbox')

            if checkbox is None:
                return None

            # Get checked state
            checked_elem = checkbox.find(qn('w14:checked'))
            if checked_elem is None:
                checked_elem = checkbox.find('{http://schemas.microsoft.com/office/word/2010/wordml}checked')

            checked = False
            if checked_elem is not None:
                val = checked_elem.get(qn('w14:val')) or checked_elem.get('{http://schemas.microsoft.com/office/word/2010/wordml}val')
                checked = val == '1' or val == 'true'

            # Get tag
            tag_elem = sdt_pr.find(qn('w:tag'))
            tag = tag_elem.get(qn('w:val')) if tag_elem is not None else ""

            # Get title/alias
            alias_elem = sdt_pr.find(qn('w:alias'))
            title = alias_elem.get(qn('w:val')) if alias_elem is not None else ""

            name = tag or title or "unnamed_checkbox"

            return CheckboxInfo(
                name=name,
                checked=checked,
                control_type="content_control",
                tag=tag,
                title=title
            )

        except Exception as e:
            logger.debug(f"Error parsing SDT checkbox: {e}")
            return None

    def _parse_legacy_checkbox(self, ffdata) -> Optional[CheckboxInfo]:
        """Parse legacy form field checkbox."""
        try:
            from docx.oxml.ns import qn

            # Check if it's a checkbox type
            checkbox_elem = ffdata.find(qn('w:checkBox'))
            if checkbox_elem is None:
                return None

            # Get name
            name_elem = ffdata.find(qn('w:name'))
            name = name_elem.get(qn('w:val')) if name_elem is not None else "unnamed"

            # Get checked state
            checked = False
            checked_elem = checkbox_elem.find(qn('w:checked'))
            if checked_elem is not None:
                checked = True
            default_elem = checkbox_elem.find(qn('w:default'))
            if default_elem is not None:
                val = default_elem.get(qn('w:val'))
                checked = val == '1' or val == 'true'

            return CheckboxInfo(
                name=name,
                checked=checked,
                control_type="form_field",
                tag=name,
                title=name
            )

        except Exception as e:
            logger.debug(f"Error parsing legacy checkbox: {e}")
            return None

    def set_checkbox(self, name: str, checked: bool) -> bool:
        """Set checkbox state by name or tag."""
        if self._docx is None:
            return False

        try:
            from docx.oxml.ns import qn

            name_lower = name.lower()
            modified = False

            # Search content control checkboxes
            for sdt in self._docx.element.iter(qn('w:sdt')):
                sdt_pr = sdt.find(qn('w:sdtPr'))
                if sdt_pr is None:
                    continue

                # Check tag and alias
                tag_elem = sdt_pr.find(qn('w:tag'))
                alias_elem = sdt_pr.find(qn('w:alias'))

                tag = (tag_elem.get(qn('w:val')) if tag_elem is not None else "").lower()
                alias = (alias_elem.get(qn('w:val')) if alias_elem is not None else "").lower()

                # Match by tag or alias
                if name_lower in tag or name_lower in alias or tag in name_lower or alias in name_lower:
                    # Find checkbox element
                    checkbox = sdt_pr.find(qn('w14:checkbox'))
                    if checkbox is None:
                        checkbox = sdt_pr.find('{http://schemas.microsoft.com/office/word/2010/wordml}checkbox')

                    if checkbox is not None:
                        # Update checked state
                        checked_elem = checkbox.find(qn('w14:checked'))
                        if checked_elem is None:
                            checked_elem = checkbox.find('{http://schemas.microsoft.com/office/word/2010/wordml}checked')

                        if checked_elem is not None:
                            nsmap = checked_elem.nsmap
                            w14_ns = nsmap.get('w14', 'http://schemas.microsoft.com/office/word/2010/wordml')
                            checked_elem.set(f'{{{w14_ns}}}val', '1' if checked else '0')
                            modified = True
                            logger.debug(f"Set checkbox '{name}' to {checked}")

            # Also handle legacy checkboxes
            for ffdata in self._docx.element.iter(qn('w:ffData')):
                name_elem = ffdata.find(qn('w:name'))
                if name_elem is None:
                    continue

                field_name = (name_elem.get(qn('w:val')) or "").lower()

                if name_lower in field_name or field_name in name_lower:
                    checkbox_elem = ffdata.find(qn('w:checkBox'))
                    if checkbox_elem is not None:
                        # Update default value
                        default_elem = checkbox_elem.find(qn('w:default'))
                        if default_elem is not None:
                            default_elem.set(qn('w:val'), '1' if checked else '0')
                            modified = True
                            logger.debug(f"Set legacy checkbox '{name}' to {checked}")

            return modified

        except Exception as e:
            logger.error(f"Error setting checkbox '{name}': {e}")
            return False

    def find_date_fields(self) -> List[FieldInfo]:
        """Find date fields in the document."""
        fields = []

        if self._docx is None:
            return fields

        try:
            from docx.oxml.ns import qn

            # Search for date content controls
            for sdt in self._docx.element.iter(qn('w:sdt')):
                sdt_pr = sdt.find(qn('w:sdtPr'))
                if sdt_pr is None:
                    continue

                # Check for date type
                date_elem = sdt_pr.find(qn('w:date'))
                if date_elem is not None:
                    tag_elem = sdt_pr.find(qn('w:tag'))
                    alias_elem = sdt_pr.find(qn('w:alias'))

                    tag = tag_elem.get(qn('w:val')) if tag_elem is not None else ""
                    title = alias_elem.get(qn('w:val')) if alias_elem is not None else ""
                    name = tag or title or "date_field"

                    # Get current value
                    sdt_content = sdt.find(qn('w:sdtContent'))
                    value = ""
                    if sdt_content is not None:
                        for t in sdt_content.iter(qn('w:t')):
                            if t.text:
                                value += t.text

                    fields.append(FieldInfo(
                        name=name,
                        field_type="content_control",
                        value=value,
                        tag=tag,
                        title=title
                    ))

                # Also check for text fields that might be date fields
                # based on tag/alias containing "fecha", "date", etc.
                tag_elem = sdt_pr.find(qn('w:tag'))
                alias_elem = sdt_pr.find(qn('w:alias'))

                tag = (tag_elem.get(qn('w:val')) if tag_elem is not None else "").lower()
                alias = (alias_elem.get(qn('w:val')) if alias_elem is not None else "").lower()

                date_keywords = ["fecha", "date", "dia", "day"]
                if any(kw in tag or kw in alias for kw in date_keywords):
                    # Already added above if it was a date control
                    if date_elem is None:
                        name = tag_elem.get(qn('w:val')) if tag_elem is not None else alias_elem.get(qn('w:val')) if alias_elem is not None else "date_field"
                        fields.append(FieldInfo(
                            name=name,
                            field_type="content_control",
                            value="",
                            tag=tag,
                            title=alias
                        ))

        except Exception as e:
            logger.warning(f"Error finding date fields: {e}")

        return fields

    def set_date_field(self, name: str, date: datetime) -> bool:
        """Set a date field value."""
        if self._docx is None:
            return False

        date_format = self.config.date_format if self.config else "%d/%m/%Y"
        date_str = date.strftime(date_format)

        try:
            from docx.oxml.ns import qn

            name_lower = name.lower()
            modified = False

            for sdt in self._docx.element.iter(qn('w:sdt')):
                sdt_pr = sdt.find(qn('w:sdtPr'))
                if sdt_pr is None:
                    continue

                tag_elem = sdt_pr.find(qn('w:tag'))
                alias_elem = sdt_pr.find(qn('w:alias'))

                tag = (tag_elem.get(qn('w:val')) if tag_elem is not None else "").lower()
                alias = (alias_elem.get(qn('w:val')) if alias_elem is not None else "").lower()

                if name_lower in tag or name_lower in alias or tag in name_lower or alias in name_lower:
                    # Update the content
                    sdt_content = sdt.find(qn('w:sdtContent'))
                    if sdt_content is not None:
                        # Find text elements and update
                        for t in sdt_content.iter(qn('w:t')):
                            t.text = date_str
                            modified = True
                            break

                        if modified:
                            logger.debug(f"Set date field '{name}' to {date_str}")
                            break

            return modified

        except Exception as e:
            logger.error(f"Error setting date field '{name}': {e}")
            return False

    def set_text_field(self, name: str, value: str) -> bool:
        """Set a text content control field value."""
        if self._docx is None:
            return False

        try:
            from docx.oxml.ns import qn

            name_lower = name.lower()

            for sdt in self._docx.element.iter(qn('w:sdt')):
                sdt_pr = sdt.find(qn('w:sdtPr'))
                if sdt_pr is None:
                    continue

                tag_elem = sdt_pr.find(qn('w:tag'))
                alias_elem = sdt_pr.find(qn('w:alias'))

                tag = (tag_elem.get(qn('w:val')) if tag_elem is not None else "").lower()
                alias = (alias_elem.get(qn('w:val')) if alias_elem is not None else "").lower()

                if name_lower in tag or name_lower in alias:
                    sdt_content = sdt.find(qn('w:sdtContent'))
                    if sdt_content is not None:
                        for t in sdt_content.iter(qn('w:t')):
                            t.text = value
                            logger.debug(f"Set text field '{name}' to '{value}'")
                            return True

            return False

        except Exception as e:
            logger.error(f"Error setting text field '{name}': {e}")
            return False


# =============================================================================
# COM Engine (Windows only)
# =============================================================================

class COMEngine(WordEngine):
    """
    Word processing engine using Windows COM automation.

    Windows only, requires Microsoft Word installed.
    Full support for all Word features including content controls.
    """

    def __init__(self, config: Any = None):
        super().__init__(config)
        self._word_app = None
        self._doc = None
        self._visible = False

    def _ensure_word_app(self) -> bool:
        """Ensure Word application is running."""
        if self._word_app is not None:
            return True

        try:
            import win32com.client
            import pythoncom

            pythoncom.CoInitialize()
            self._word_app = win32com.client.Dispatch("Word.Application")
            self._word_app.Visible = self._visible
            self._word_app.DisplayAlerts = False

            logger.debug("Word COM application started")
            return True

        except ImportError:
            logger.error("pywin32 not installed. Install with: pip install pywin32")
            return False
        except Exception as e:
            logger.error(f"Failed to start Word application: {e}")
            return False

    def open(self, file_path: Path) -> bool:
        """Open a Word document using COM."""
        if not self._ensure_word_app():
            return False

        try:
            self._document_path = Path(file_path).resolve()
            self._doc = self._word_app.Documents.Open(str(self._document_path))
            self._document = self._doc

            logger.debug(f"Opened document with COM: {file_path}")
            return True

        except Exception as e:
            logger.error(f"Failed to open document with COM: {e}")
            return False

    def save(self, output_path: Optional[Path] = None) -> bool:
        """Save the document."""
        if self._doc is None:
            logger.error("No document open to save")
            return False

        try:
            if output_path:
                output_path = Path(output_path).resolve()
                # SaveAs2 with FileFormat 16 = docx
                self._doc.SaveAs2(str(output_path), FileFormat=16)
            else:
                self._doc.Save()

            logger.debug(f"Saved document: {output_path or self._document_path}")
            return True

        except Exception as e:
            logger.error(f"Failed to save document: {e}")
            return False

    def close(self) -> None:
        """Close the document and Word application."""
        try:
            if self._doc is not None:
                self._doc.Close(SaveChanges=False)
                self._doc = None
                self._document = None
        except Exception as e:
            logger.warning(f"Error closing document: {e}")

        try:
            if self._word_app is not None:
                self._word_app.Quit()
                self._word_app = None
        except Exception as e:
            logger.warning(f"Error closing Word: {e}")

        try:
            import pythoncom
            pythoncom.CoUninitialize()
        except Exception:
            pass

        self._document_path = None

    def get_text_content(self) -> str:
        """Get document text content."""
        if self._doc is None:
            return ""

        try:
            return self._doc.Content.Text
        except Exception as e:
            logger.error(f"Error getting text content: {e}")
            return ""

    def find_checkboxes(self) -> List[CheckboxInfo]:
        """Find all checkboxes in the document."""
        checkboxes = []

        if self._doc is None:
            return checkboxes

        try:
            # Search Content Controls (newer style)
            for cc in self._doc.ContentControls:
                # Type 8 = wdContentControlCheckBox
                if cc.Type == 8:
                    checkboxes.append(CheckboxInfo(
                        name=cc.Tag or cc.Title or f"checkbox_{cc.ID}",
                        checked=cc.Checked,
                        control_type="content_control",
                        tag=cc.Tag or "",
                        title=cc.Title or ""
                    ))

            # Search legacy form fields
            for field in self._doc.FormFields:
                # Type 71 = wdFieldFormCheckBox
                if field.Type == 71:
                    checkboxes.append(CheckboxInfo(
                        name=field.Name,
                        checked=field.CheckBox.Value,
                        control_type="form_field",
                        tag=field.Name,
                        title=field.Name
                    ))

        except Exception as e:
            logger.warning(f"Error finding checkboxes: {e}")

        return checkboxes

    def set_checkbox(self, name: str, checked: bool) -> bool:
        """Set checkbox state by name or tag."""
        if self._doc is None:
            return False

        name_lower = name.lower()
        modified = False

        try:
            # Search Content Controls
            for cc in self._doc.ContentControls:
                if cc.Type == 8:  # Checkbox
                    tag = (cc.Tag or "").lower()
                    title = (cc.Title or "").lower()

                    if (name_lower in tag or name_lower in title or
                        tag in name_lower or title in name_lower):
                        cc.Checked = checked
                        modified = True
                        logger.debug(f"Set content control checkbox '{name}' to {checked}")
                        break

            # Search legacy form fields if not found
            if not modified:
                for field in self._doc.FormFields:
                    if field.Type == 71:  # Checkbox
                        field_name = (field.Name or "").lower()

                        if name_lower in field_name or field_name in name_lower:
                            field.CheckBox.Value = checked
                            modified = True
                            logger.debug(f"Set form field checkbox '{name}' to {checked}")
                            break

        except Exception as e:
            logger.error(f"Error setting checkbox '{name}': {e}")

        return modified

    def find_date_fields(self) -> List[FieldInfo]:
        """Find date fields in the document."""
        fields = []

        if self._doc is None:
            return fields

        try:
            date_keywords = ["fecha", "date", "dia", "day"]

            for cc in self._doc.ContentControls:
                tag = (cc.Tag or "").lower()
                title = (cc.Title or "").lower()

                # Check if it's a date picker (Type 6) or text with date keyword
                is_date = cc.Type == 6  # wdContentControlDate
                is_date = is_date or any(kw in tag or kw in title for kw in date_keywords)

                if is_date:
                    fields.append(FieldInfo(
                        name=cc.Tag or cc.Title or f"date_{cc.ID}",
                        field_type="content_control",
                        value=cc.Range.Text if hasattr(cc, 'Range') else "",
                        tag=cc.Tag or "",
                        title=cc.Title or ""
                    ))

        except Exception as e:
            logger.warning(f"Error finding date fields: {e}")

        return fields

    def set_date_field(self, name: str, date: datetime) -> bool:
        """Set a date field value."""
        if self._doc is None:
            return False

        date_format = self.config.date_format if self.config else "%d/%m/%Y"
        date_str = date.strftime(date_format)
        name_lower = name.lower()

        try:
            for cc in self._doc.ContentControls:
                tag = (cc.Tag or "").lower()
                title = (cc.Title or "").lower()

                if name_lower in tag or name_lower in title or tag in name_lower or title in name_lower:
                    # For date picker controls
                    if cc.Type == 6:  # wdContentControlDate
                        cc.Range.Text = date_str
                        logger.debug(f"Set date picker '{name}' to {date_str}")
                        return True
                    # For text controls
                    elif cc.Type in (1, 0):  # wdContentControlRichText or wdContentControlText
                        cc.Range.Text = date_str
                        logger.debug(f"Set text field '{name}' to {date_str}")
                        return True

        except Exception as e:
            logger.error(f"Error setting date field '{name}': {e}")

        return False

    def set_text_field(self, name: str, value: str) -> bool:
        """Set a text content control field value."""
        if self._doc is None:
            return False

        name_lower = name.lower()

        try:
            for cc in self._doc.ContentControls:
                tag = (cc.Tag or "").lower()
                title = (cc.Title or "").lower()

                if name_lower in tag or name_lower in title:
                    cc.Range.Text = value
                    logger.debug(f"Set text field '{name}' to '{value}'")
                    return True

        except Exception as e:
            logger.error(f"Error setting text field '{name}': {e}")

        return False


# =============================================================================
# Engine Factory
# =============================================================================

def get_engine(engine_type: str = "docx", config: Any = None) -> WordEngine:
    """
    Get a Word processing engine instance.

    Args:
        engine_type: "docx" or "com"
        config: WordConfig object

    Returns:
        WordEngine instance

    Raises:
        ValueError: If engine type is invalid
        RuntimeError: If COM engine requested on non-Windows
    """
    engine_type = engine_type.lower()

    if engine_type == "docx":
        return DocxEngine(config)
    elif engine_type == "com":
        import sys
        if sys.platform != "win32":
            raise RuntimeError("COM engine requires Windows")
        return COMEngine(config)
    else:
        raise ValueError(f"Unknown engine type: {engine_type}")


# =============================================================================
# Document Processor
# =============================================================================

class DocumentProcessor:
    """
    High-level document processor that handles the full workflow.

    Coordinates opening, modifying, and saving documents using the
    configured engine.
    """

    def __init__(self, config: Any, rules: Any = None):
        """
        Initialize the processor.

        Args:
            config: Config object with all settings
            rules: RulesConfig object with processing rules
        """
        self.config = config
        self.rules = rules
        self._engine: Optional[WordEngine] = None

    def _get_engine(self) -> WordEngine:
        """Get or create the Word engine."""
        if self._engine is None:
            engine_type = self.config.word.engine if self.config else "docx"

            # Fall back to docx engine on non-Windows
            import sys
            if engine_type == "com" and sys.platform != "win32":
                logger.warning("COM engine not available on non-Windows, falling back to docx")
                engine_type = "docx"

            self._engine = get_engine(engine_type, self.config.word if self.config else None)

        return self._engine

    def process_document(
        self,
        input_path: Path,
        output_path: Optional[Path] = None,
        dry_run: bool = False,
        set_date: bool = True,
        checkbox_overrides: Optional[Dict[str, bool]] = None,
    ) -> ProcessingResult:
        """
        Process a single document.

        Args:
            input_path: Path to input .docx file
            output_path: Path for output file (default: processed/docx/)
            dry_run: If True, don't save changes
            set_date: If True, set date fields to today
            checkbox_overrides: Dict of checkbox_name -> checked state

        Returns:
            ProcessingResult with details of the operation
        """
        input_path = Path(input_path)
        result = ProcessingResult(
            success=False,
            input_path=input_path
        )

        # Determine output path
        if output_path is None and self.config:
            output_dir = self.config.paths.output_docx
            output_path = output_dir / input_path.name
            # Ensure unique filename
            output_path = ensure_unique_filename(output_path)

        result.output_path = output_path

        # Dry run check
        if dry_run:
            logger.info(f"[DRY RUN] Would process: {input_path}")
            result.success = True
            return result

        engine = self._get_engine()

        try:
            # Open document
            if not engine.open(input_path):
                result.error_message = "Failed to open document"
                return result

            # Get document text for rule evaluation
            doc_text = engine.get_text_content()

            # Set date fields
            if set_date:
                date_fields = self._get_date_field_names()
                today = datetime.now()

                for field_name in date_fields:
                    if engine.set_date_field(field_name, today):
                        result.fields_modified.append(field_name)
                        logger.info(f"Set date field '{field_name}' to {today.strftime('%d/%m/%Y')}")

            # Apply checkbox overrides
            if checkbox_overrides:
                for name, checked in checkbox_overrides.items():
                    if engine.set_checkbox(name, checked):
                        result.checkboxes_modified.append(name)
                        logger.info(f"Set checkbox '{name}' to {checked}")

            # Apply rules-based checkboxes
            if self.rules:
                checkbox_actions = self._evaluate_rules(doc_text, input_path.name)
                for name, checked in checkbox_actions.items():
                    if name not in (checkbox_overrides or {}):
                        if engine.set_checkbox(name, checked):
                            result.checkboxes_modified.append(name)
                            logger.info(f"Set checkbox '{name}' to {checked} (by rule)")

            # Save document
            if output_path:
                output_path.parent.mkdir(parents=True, exist_ok=True)

            if engine.save(output_path):
                result.success = True
                logger.info(f"Saved processed document: {output_path}")
            else:
                result.error_message = "Failed to save document"

        except Exception as e:
            result.error_message = str(e)
            logger.error(f"Error processing document: {e}")

        finally:
            engine.close()

        return result

    def _get_date_field_names(self) -> List[str]:
        """Get list of date field names to look for."""
        if self.config and self.config.word.date_field_names:
            return self.config.word.date_field_names
        return ["Fecha", "DATE", "FechaDocumento"]

    def _evaluate_rules(self, doc_text: str, filename: str) -> Dict[str, bool]:
        """
        Evaluate rules against document content.

        Args:
            doc_text: Document text content
            filename: Document filename

        Returns:
            Dict of checkbox_name -> checked state
        """
        checkbox_actions = {}

        if not self.rules:
            return checkbox_actions

        doc_text_lower = doc_text.lower()
        filename_lower = filename.lower()

        for rule in self.rules.get_enabled_rules():
            if not self._check_conditions(rule.conditions, doc_text_lower, filename_lower):
                continue

            # Apply actions
            for action in rule.actions:
                action_type = action.get("type", "")

                if action_type == "set_checkbox":
                    checkbox_name = action.get("checkbox", "")
                    checked = action.get("value", True)

                    # Use mappings to find actual checkbox names
                    possible_names = self.rules.find_checkbox_tags(checkbox_name)
                    for name in possible_names:
                        checkbox_actions[name] = checked

        return checkbox_actions

    def _check_conditions(
        self,
        conditions: List[Dict[str, Any]],
        doc_text: str,
        filename: str
    ) -> bool:
        """Check if all conditions are met."""
        if not conditions:
            return True  # No conditions = always applies

        # Normalize to lowercase for case-insensitive matching
        doc_text_lower = doc_text.lower()
        filename_lower = filename.lower()

        for condition in conditions:
            cond_type = condition.get("type", "")

            if cond_type == "contains":
                value = condition.get("value", "").lower()
                if value not in doc_text_lower:
                    return False

            elif cond_type == "not_contains":
                value = condition.get("value", "").lower()
                if value in doc_text_lower:
                    return False

            elif cond_type == "contains_any":
                values = [v.lower() for v in condition.get("values", [])]
                if not any(v in doc_text_lower for v in values):
                    return False

            elif cond_type == "contains_all":
                values = [v.lower() for v in condition.get("values", [])]
                if not all(v in doc_text_lower for v in values):
                    return False

            elif cond_type == "filename_contains":
                value = condition.get("value", "").lower()
                if value not in filename_lower:
                    return False

            elif cond_type == "filename_regex":
                pattern = condition.get("pattern", "")
                if not re.search(pattern, filename, re.IGNORECASE):
                    return False

            elif cond_type == "regex":
                pattern = condition.get("pattern", "")
                if not re.search(pattern, doc_text, re.IGNORECASE):
                    return False

        return True
