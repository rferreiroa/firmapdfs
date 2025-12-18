"""
Tests for Word Processor Module
===============================
"""

import pytest
from pathlib import Path
from datetime import datetime
from unittest.mock import Mock, patch, MagicMock

from firmapdfs.word_processor import (
    DocxEngine,
    COMEngine,
    DocumentProcessor,
    ProcessingResult,
    CheckboxInfo,
    FieldInfo,
    get_engine,
)


# =============================================================================
# Test Fixtures
# =============================================================================

@pytest.fixture
def mock_config():
    """Create mock configuration."""
    config = Mock()
    config.word = Mock()
    config.word.engine = "docx"
    config.word.date_format = "%d/%m/%Y"
    config.word.date_field_names = ["Fecha", "DATE"]
    config.paths = Mock()
    config.paths.output_docx = Path("/tmp/output/docx")
    return config


@pytest.fixture
def mock_rules():
    """Create mock rules configuration."""
    rules = Mock()
    rules.get_enabled_rules.return_value = []
    rules.find_checkbox_tags.return_value = ["chkTest"]
    return rules


@pytest.fixture
def sample_docx(tmp_path):
    """Create a sample DOCX file for testing."""
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        doc = Document()
        doc.add_paragraph("Test Document")
        doc.add_paragraph("This is a test document for FirmaPDFs.")
        doc.add_paragraph("Contains: APROBADO keyword for rule testing.")

        # Save the document
        docx_path = tmp_path / "test_document.docx"
        doc.save(str(docx_path))

        return docx_path

    except ImportError:
        pytest.skip("python-docx not installed")


@pytest.fixture
def sample_docx_with_controls(tmp_path):
    """Create a DOCX with content controls (if possible)."""
    try:
        from docx import Document

        doc = Document()
        doc.add_paragraph("Document with Controls")
        doc.add_paragraph("Fecha: [DATE FIELD]")
        doc.add_paragraph("Estado: [CHECKBOX]")

        docx_path = tmp_path / "test_controls.docx"
        doc.save(str(docx_path))

        return docx_path

    except ImportError:
        pytest.skip("python-docx not installed")


# =============================================================================
# Engine Factory Tests
# =============================================================================

class TestEngineFactory:
    """Tests for engine factory function."""

    def test_get_docx_engine(self):
        """Test getting docx engine."""
        engine = get_engine("docx")
        assert isinstance(engine, DocxEngine)

    def test_get_docx_engine_case_insensitive(self):
        """Test engine selection is case insensitive."""
        engine = get_engine("DOCX")
        assert isinstance(engine, DocxEngine)

    def test_invalid_engine_type(self):
        """Test invalid engine type raises error."""
        with pytest.raises(ValueError, match="Unknown engine type"):
            get_engine("invalid")

    @pytest.mark.skipif(
        __import__("sys").platform != "win32",
        reason="COM engine only available on Windows"
    )
    def test_get_com_engine_windows(self):
        """Test getting COM engine on Windows."""
        engine = get_engine("com")
        assert isinstance(engine, COMEngine)

    @pytest.mark.skipif(
        __import__("sys").platform == "win32",
        reason="Test for non-Windows systems"
    )
    def test_com_engine_non_windows(self):
        """Test COM engine raises error on non-Windows."""
        with pytest.raises(RuntimeError, match="requires Windows"):
            get_engine("com")


# =============================================================================
# DocxEngine Tests
# =============================================================================

class TestDocxEngine:
    """Tests for DocxEngine."""

    def test_open_nonexistent_file(self, mock_config):
        """Test opening non-existent file returns False."""
        engine = DocxEngine(mock_config)
        result = engine.open(Path("/nonexistent/file.docx"))
        assert result is False

    def test_open_valid_file(self, mock_config, sample_docx):
        """Test opening a valid DOCX file."""
        engine = DocxEngine(mock_config)
        result = engine.open(sample_docx)

        assert result is True
        engine.close()

    def test_get_text_content(self, mock_config, sample_docx):
        """Test getting text content from document."""
        engine = DocxEngine(mock_config)
        engine.open(sample_docx)

        text = engine.get_text_content()

        assert "Test Document" in text
        assert "APROBADO" in text

        engine.close()

    def test_save_document(self, mock_config, sample_docx, tmp_path):
        """Test saving document to new location."""
        engine = DocxEngine(mock_config)
        engine.open(sample_docx)

        output_path = tmp_path / "output" / "saved.docx"
        output_path.parent.mkdir(parents=True, exist_ok=True)

        result = engine.save(output_path)

        assert result is True
        assert output_path.exists()

        engine.close()

    def test_save_without_open(self, mock_config):
        """Test saving without opening returns False."""
        engine = DocxEngine(mock_config)
        result = engine.save(Path("/tmp/test.docx"))
        assert result is False

    def test_context_manager(self, mock_config, sample_docx):
        """Test using engine as context manager."""
        with DocxEngine(mock_config) as engine:
            result = engine.open(sample_docx)
            assert result is True
            text = engine.get_text_content()
            assert "Test Document" in text

    def test_find_checkboxes_empty_doc(self, mock_config, sample_docx):
        """Test finding checkboxes in document without controls."""
        engine = DocxEngine(mock_config)
        engine.open(sample_docx)

        checkboxes = engine.find_checkboxes()

        # Simple document shouldn't have checkboxes
        assert isinstance(checkboxes, list)

        engine.close()

    def test_find_date_fields_empty_doc(self, mock_config, sample_docx):
        """Test finding date fields in document without controls."""
        engine = DocxEngine(mock_config)
        engine.open(sample_docx)

        fields = engine.find_date_fields()

        assert isinstance(fields, list)

        engine.close()


# =============================================================================
# DocumentProcessor Tests
# =============================================================================

class TestDocumentProcessor:
    """Tests for DocumentProcessor."""

    def test_processor_initialization(self, mock_config, mock_rules):
        """Test processor initialization."""
        processor = DocumentProcessor(mock_config, mock_rules)

        assert processor.config == mock_config
        assert processor.rules == mock_rules

    def test_process_nonexistent_file(self, mock_config, mock_rules):
        """Test processing non-existent file."""
        processor = DocumentProcessor(mock_config, mock_rules)

        result = processor.process_document(
            input_path=Path("/nonexistent/file.docx"),
            dry_run=False
        )

        assert result.success is False
        assert "open" in result.error_message.lower() or result.error_message != ""

    def test_process_dry_run(self, mock_config, mock_rules, sample_docx):
        """Test processing in dry run mode."""
        processor = DocumentProcessor(mock_config, mock_rules)

        result = processor.process_document(
            input_path=sample_docx,
            dry_run=True
        )

        assert result.success is True
        assert result.input_path == sample_docx

    def test_process_document_success(self, mock_config, mock_rules, sample_docx, tmp_path):
        """Test successful document processing."""
        # Setup output directory
        output_dir = tmp_path / "output" / "docx"
        output_dir.mkdir(parents=True, exist_ok=True)
        mock_config.paths.output_docx = output_dir

        processor = DocumentProcessor(mock_config, mock_rules)

        result = processor.process_document(
            input_path=sample_docx,
            dry_run=False
        )

        assert result.success is True
        assert result.output_path is not None
        assert result.output_path.exists()

    def test_process_with_checkbox_overrides(self, mock_config, mock_rules, sample_docx, tmp_path):
        """Test processing with checkbox overrides."""
        output_dir = tmp_path / "output" / "docx"
        output_dir.mkdir(parents=True, exist_ok=True)
        mock_config.paths.output_docx = output_dir

        processor = DocumentProcessor(mock_config, mock_rules)

        result = processor.process_document(
            input_path=sample_docx,
            checkbox_overrides={"aprobado": True, "rechazado": False}
        )

        assert result.success is True


# =============================================================================
# ProcessingResult Tests
# =============================================================================

class TestProcessingResult:
    """Tests for ProcessingResult dataclass."""

    def test_default_values(self):
        """Test default values."""
        result = ProcessingResult(
            success=True,
            input_path=Path("/test/file.docx")
        )

        assert result.success is True
        assert result.output_path is None
        assert result.error_message == ""
        assert result.fields_modified == []
        assert result.checkboxes_modified == []

    def test_with_modifications(self):
        """Test result with modifications."""
        result = ProcessingResult(
            success=True,
            input_path=Path("/test/file.docx"),
            output_path=Path("/output/file.docx"),
            fields_modified=["Fecha"],
            checkboxes_modified=["aprobado", "conforme"]
        )

        assert len(result.fields_modified) == 1
        assert len(result.checkboxes_modified) == 2


# =============================================================================
# CheckboxInfo Tests
# =============================================================================

class TestCheckboxInfo:
    """Tests for CheckboxInfo dataclass."""

    def test_creation(self):
        """Test checkbox info creation."""
        info = CheckboxInfo(
            name="chkAprobado",
            checked=True,
            control_type="content_control",
            tag="aprobado",
            title="Aprobado"
        )

        assert info.name == "chkAprobado"
        assert info.checked is True
        assert info.control_type == "content_control"


# =============================================================================
# FieldInfo Tests
# =============================================================================

class TestFieldInfo:
    """Tests for FieldInfo dataclass."""

    def test_creation(self):
        """Test field info creation."""
        info = FieldInfo(
            name="Fecha",
            field_type="content_control",
            value="15/01/2024",
            tag="fecha",
            title="Fecha del documento"
        )

        assert info.name == "Fecha"
        assert info.field_type == "content_control"
        assert info.value == "15/01/2024"


# =============================================================================
# Rule Evaluation Tests
# =============================================================================

class TestRuleEvaluation:
    """Tests for rule evaluation in DocumentProcessor."""

    def test_condition_contains(self, mock_config):
        """Test 'contains' condition."""
        mock_rules = Mock()
        rule = Mock()
        rule.conditions = [{"type": "contains", "value": "APROBADO"}]
        rule.actions = [{"type": "set_checkbox", "checkbox": "aprobado", "value": True}]
        mock_rules.get_enabled_rules.return_value = [rule]
        mock_rules.find_checkbox_tags.return_value = ["chkAprobado"]

        processor = DocumentProcessor(mock_config, mock_rules)

        # Test _check_conditions directly
        result = processor._check_conditions(
            rule.conditions,
            "Este documento está APROBADO",
            "test.docx"
        )

        assert result is True

    def test_condition_contains_fail(self, mock_config):
        """Test 'contains' condition that fails."""
        processor = DocumentProcessor(mock_config, None)

        conditions = [{"type": "contains", "value": "aprobado"}]
        result = processor._check_conditions(
            conditions,
            "documento rechazado",
            "test.docx"
        )

        assert result is False

    def test_condition_contains_any(self, mock_config):
        """Test 'contains_any' condition."""
        processor = DocumentProcessor(mock_config, None)

        conditions = [{
            "type": "contains_any",
            "values": ["aprobado", "aceptado", "ok"]
        }]

        # Should match "aprobado"
        result = processor._check_conditions(
            conditions,
            "documento aprobado",
            "test.docx"
        )
        assert result is True

        # Should fail
        result = processor._check_conditions(
            conditions,
            "documento pendiente",
            "test.docx"
        )
        assert result is False

    def test_condition_filename_contains(self, mock_config):
        """Test 'filename_contains' condition."""
        processor = DocumentProcessor(mock_config, None)

        conditions = [{"type": "filename_contains", "value": "urgente"}]

        result = processor._check_conditions(
            conditions,
            "contenido del documento",
            "URGENTE_informe.docx"
        )

        assert result is True

    def test_condition_empty(self, mock_config):
        """Test empty conditions always pass."""
        processor = DocumentProcessor(mock_config, None)

        result = processor._check_conditions([], "any text", "any.docx")

        assert result is True

    def test_evaluate_rules(self, mock_config):
        """Test full rule evaluation."""
        mock_rules = Mock()

        rule1 = Mock()
        rule1.conditions = [{"type": "contains", "value": "aprobado"}]
        rule1.actions = [
            {"type": "set_checkbox", "checkbox": "aprobado", "value": True}
        ]

        mock_rules.get_enabled_rules.return_value = [rule1]
        mock_rules.find_checkbox_tags.return_value = ["chkAprobado", "Aprobado"]

        processor = DocumentProcessor(mock_config, mock_rules)

        actions = processor._evaluate_rules(
            "El documento está APROBADO",
            "test.docx"
        )

        assert "chkAprobado" in actions or "Aprobado" in actions
