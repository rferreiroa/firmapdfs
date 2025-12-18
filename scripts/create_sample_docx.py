#!/usr/bin/env python
"""
Script to create sample Word documents for testing FirmaPDFs.

This script generates test documents that can be used to verify
FirmaPDFs functionality.

Usage:
    python scripts/create_sample_docx.py
"""

import sys
from pathlib import Path
from lxml import etree

# Add src to path
sys.path.insert(0, str(Path(__file__).parent.parent / "src"))

try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import register_element_cls
    from lxml.builder import ElementMaker
except ImportError as e:
    print(f"Error: Missing dependency. Run: pip install python-docx lxml")
    print(f"Details: {e}")
    sys.exit(1)


# Word 2010+ namespace for checkboxes
W14_NS = "http://schemas.microsoft.com/office/word/2010/wordml"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def create_checkbox_sdt(tag_name: str, title: str, checked: bool = False):
    """
    Create an SDT (Structured Document Tag) checkbox element.

    This creates the XML structure for a Word checkbox content control.
    """
    nsmap = {
        'w': W_NS,
        'w14': W14_NS,
    }

    # Create SDT element
    sdt = etree.Element(qn('w:sdt'), nsmap=nsmap)

    # SDT Properties
    sdt_pr = etree.SubElement(sdt, qn('w:sdtPr'))

    # Tag
    tag_elem = etree.SubElement(sdt_pr, qn('w:tag'))
    tag_elem.set(qn('w:val'), tag_name)

    # Alias/Title
    alias = etree.SubElement(sdt_pr, qn('w:alias'))
    alias.set(qn('w:val'), title)

    # ID
    sdt_id = etree.SubElement(sdt_pr, qn('w:id'))
    sdt_id.set(qn('w:val'), str(hash(tag_name) % 100000))

    # Checkbox element with w14 namespace
    checkbox = etree.SubElement(sdt_pr, '{%s}checkbox' % W14_NS)

    checked_elem = etree.SubElement(checkbox, '{%s}checked' % W14_NS)
    checked_elem.set('{%s}val' % W14_NS, '1' if checked else '0')

    checked_state = etree.SubElement(checkbox, '{%s}checkedState' % W14_NS)
    checked_state.set('{%s}val' % W14_NS, '2612')
    checked_state.set('{%s}font' % W14_NS, 'MS Gothic')

    unchecked_state = etree.SubElement(checkbox, '{%s}uncheckedState' % W14_NS)
    unchecked_state.set('{%s}val' % W14_NS, '2610')
    unchecked_state.set('{%s}font' % W14_NS, 'MS Gothic')

    # SDT Content
    sdt_content = etree.SubElement(sdt, qn('w:sdtContent'))

    # Run with the checkbox character
    run = etree.SubElement(sdt_content, qn('w:r'))
    run_pr = etree.SubElement(run, qn('w:rPr'))
    fonts = etree.SubElement(run_pr, qn('w:rFonts'))
    fonts.set(qn('w:ascii'), 'MS Gothic')
    fonts.set(qn('w:hAnsi'), 'MS Gothic')

    text = etree.SubElement(run, qn('w:t'))
    text.text = '☒' if checked else '☐'

    return sdt


def create_text_sdt(tag_name: str, title: str, placeholder: str = ""):
    """Create a text content control SDT element."""
    nsmap = {
        'w': W_NS,
    }

    sdt = etree.Element(qn('w:sdt'), nsmap=nsmap)

    # SDT Properties
    sdt_pr = etree.SubElement(sdt, qn('w:sdtPr'))

    # Tag
    tag_elem = etree.SubElement(sdt_pr, qn('w:tag'))
    tag_elem.set(qn('w:val'), tag_name)

    # Alias
    alias = etree.SubElement(sdt_pr, qn('w:alias'))
    alias.set(qn('w:val'), title)

    # ID
    sdt_id = etree.SubElement(sdt_pr, qn('w:id'))
    sdt_id.set(qn('w:val'), str(hash(tag_name) % 100000))

    # Text type
    etree.SubElement(sdt_pr, qn('w:text'))

    # SDT Content
    sdt_content = etree.SubElement(sdt, qn('w:sdtContent'))

    run = etree.SubElement(sdt_content, qn('w:r'))
    text = etree.SubElement(run, qn('w:t'))
    text.text = placeholder or "[Click to edit]"

    return sdt


def add_checkbox_to_paragraph(paragraph, tag_name: str, title: str, checked: bool = False):
    """Add a checkbox content control to a paragraph."""
    sdt = create_checkbox_sdt(tag_name, title, checked)
    paragraph._p.append(sdt)


def add_text_control_to_paragraph(paragraph, tag_name: str, title: str, placeholder: str = ""):
    """Add a text content control to a paragraph."""
    sdt = create_text_sdt(tag_name, title, placeholder)
    paragraph._p.append(sdt)


def create_approval_document(output_path: Path):
    """Create an approval document template with content controls."""
    doc = Document()

    # Title
    title = doc.add_heading("DOCUMENTO DE APROBACIÓN", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # Date field
    p = doc.add_paragraph()
    p.add_run("Fecha: ")
    add_text_control_to_paragraph(p, "Fecha", "Fecha del documento", "__/__/____")

    # Reference
    p = doc.add_paragraph()
    p.add_run("Referencia: ")
    add_text_control_to_paragraph(p, "Referencia", "Número de referencia", "REF-2024-XXX")

    doc.add_paragraph()

    # Content
    doc.add_heading("Contenido", level=2)
    doc.add_paragraph(
        "Este documento ha sido revisado y se encuentra APROBADO según los "
        "criterios establecidos en el procedimiento interno de calidad."
    )
    doc.add_paragraph(
        "El documento cumple con todos los requisitos necesarios y "
        "la documentación está completa."
    )

    doc.add_paragraph()

    # Resolution section
    doc.add_heading("Resolución", level=2)

    # Checkboxes
    p = doc.add_paragraph()
    add_checkbox_to_paragraph(p, "chkAprobado", "Aprobado", False)
    p.add_run(" Aprobado")

    p = doc.add_paragraph()
    add_checkbox_to_paragraph(p, "chkRechazado", "Rechazado", False)
    p.add_run(" Rechazado")

    p = doc.add_paragraph()
    add_checkbox_to_paragraph(p, "chkPendiente", "Pendiente", False)
    p.add_run(" Pendiente de revisión")

    doc.add_paragraph()

    # Conformity
    doc.add_heading("Conformidad", level=2)

    p = doc.add_paragraph()
    add_checkbox_to_paragraph(p, "chkConforme", "Conforme", False)
    p.add_run(" Conforme")

    p = doc.add_paragraph()
    add_checkbox_to_paragraph(p, "chkNoConforme", "No Conforme", False)
    p.add_run(" No conforme")

    doc.add_paragraph()

    # Signature section
    doc.add_heading("Firma", level=2)

    p = doc.add_paragraph()
    p.add_run("Responsable: ")
    add_text_control_to_paragraph(p, "Responsable", "Nombre del responsable", "[Nombre]")

    p = doc.add_paragraph()
    p.add_run("Departamento: ")
    add_text_control_to_paragraph(p, "Departamento", "Departamento", "[Departamento]")

    doc.add_paragraph()
    doc.add_paragraph("_" * 40)
    doc.add_paragraph("Firma")

    # Save
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Created: {output_path}")


def create_inspection_document(output_path: Path):
    """Create an inspection report template."""
    doc = Document()

    title = doc.add_heading("INFORME DE INSPECCIÓN", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # Header
    p = doc.add_paragraph()
    p.add_run("Fecha: ")
    add_text_control_to_paragraph(p, "Fecha", "Fecha de inspección", "__/__/____")

    p = doc.add_paragraph()
    p.add_run("Inspector: ")
    add_text_control_to_paragraph(p, "Inspector", "Nombre", "[Inspector]")

    doc.add_paragraph()

    # Results
    doc.add_heading("Resultados", level=2)
    doc.add_paragraph("Inspección realizada conforme al procedimiento establecido.")

    # Checklist
    doc.add_heading("Lista de Verificación", level=2)

    items = [
        ("chkDocumentacion", "Documentación completa"),
        ("chkSeguridad", "Requisitos de seguridad"),
        ("chkCalidad", "Estándares de calidad"),
    ]

    for tag, label in items:
        p = doc.add_paragraph()
        add_checkbox_to_paragraph(p, tag, label, False)
        p.add_run(f" {label}")

    doc.add_paragraph()

    # Final verdict
    doc.add_heading("Dictamen", level=2)

    p = doc.add_paragraph()
    add_checkbox_to_paragraph(p, "chkAprobado", "Aprobado", False)
    p.add_run(" APROBADO")

    p = doc.add_paragraph()
    add_checkbox_to_paragraph(p, "chkRechazado", "Rechazado", False)
    p.add_run(" RECHAZADO")

    doc.add_paragraph()
    doc.add_paragraph("_" * 40)
    doc.add_paragraph("Firma del Inspector")

    # Save
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Created: {output_path}")


def create_simple_document(output_path: Path):
    """Create a simple document without content controls."""
    doc = Document()

    doc.add_heading("Documento Simple de Prueba", level=1)
    doc.add_paragraph("Fecha: [FECHA]")
    doc.add_paragraph()
    doc.add_paragraph(
        "Este es un documento simple sin controles de contenido. "
        "Se puede usar para pruebas básicas del sistema."
    )
    doc.add_paragraph()
    doc.add_paragraph("El documento está APROBADO y es CONFORME con los requisitos.")
    doc.add_paragraph()
    doc.add_paragraph("Firma: ____________________")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f"Created: {output_path}")


def main():
    """Generate all sample documents."""
    base_path = Path(__file__).parent.parent

    templates_dir = base_path / "samples" / "templates"
    inbox_dir = base_path / "samples" / "inbox"

    print("Creating sample Word documents...")
    print()

    # Create templates
    create_approval_document(templates_dir / "plantilla_aprobacion.docx")
    create_inspection_document(templates_dir / "plantilla_inspeccion.docx")
    create_simple_document(templates_dir / "plantilla_simple.docx")

    # Create test files in inbox
    create_approval_document(inbox_dir / "documento_aprobacion_001.docx")
    create_simple_document(inbox_dir / "documento_simple_001.docx")

    print()
    print("Sample documents created successfully!")
    print()
    print(f"Templates: {templates_dir}")
    print(f"Test files: {inbox_dir}")


if __name__ == "__main__":
    main()
